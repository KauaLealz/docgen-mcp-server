"""Conversao de HTML simples para apresentacoes PowerPoint."""

import base64
import os
import re
import tempfile
import zipfile
from io import BytesIO

from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from security import validate_path, validate_write_path


SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5
MARGIN = 0.5
CONTENT_TOP = 1.2
TITLE_HEIGHT = 0.8


class PptxRenderError(ValueError):
    """Erro de renderizacao com contexto suficiente para corrigir a entrada."""


def create_pptx_from_html(
    html: str,
    output_path: str,
    title: str | None = None,
    author: str | None = None,
    aspect_ratio: str = "16:9",
    css: str = "",
    footer_html: str | None = None,
    asset_base_path: str | None = None,
) -> str:
    """Gera um PPTX a partir de HTML, usando um slide por section/.slide."""
    out_path = validate_write_path(output_path)
    if not out_path.lower().endswith(".pptx"):
        out_path = os.path.splitext(out_path)[0] + ".pptx"

    if aspect_ratio not in {"16:9", "4:3"}:
        raise ValueError("render_slide: aspectRatio deve ser '16:9' ou '4:3'.")
    width, height = (SLIDE_WIDTH, SLIDE_HEIGHT) if aspect_ratio == "16:9" else (10, 7.5)
    soup = BeautifulSoup(html, "html.parser")
    _apply_basic_css(soup, css)
    blocks = _extract_slide_blocks(soup) or [soup]

    presentation = Presentation()
    presentation.slide_width = Inches(width)
    presentation.slide_height = Inches(height)
    blank = presentation.slide_layouts[6]
    for index, block in enumerate(blocks):
        slide = presentation.slides.add_slide(blank)
        try:
            _render_slide(slide, block, title if index == 0 else None, width, height,
                          index + 1, len(blocks), footer_html, asset_base_path)
        except Exception as exc:
            if isinstance(exc, PptxRenderError):
                raise PptxRenderError(f"slide {index + 1}/{len(blocks)}: {exc}") from exc
            raise PptxRenderError(f"slide {index + 1}/{len(blocks)} falhou: {exc}") from exc

    if title:
        presentation.core_properties.title = title
    if author:
        presentation.core_properties.author = author
    output_dir = os.path.dirname(out_path) or os.getcwd()
    os.makedirs(output_dir, exist_ok=True)
    fd, temp_path = tempfile.mkstemp(prefix=".docgen-pptx-", suffix=".tmp", dir=output_dir)
    os.close(fd)
    try:
        presentation.save(temp_path)
        _validate_pptx(temp_path)
        os.replace(temp_path, out_path)
    finally:
        try:
            os.unlink(temp_path)
        except OSError:
            pass
    return out_path


def _extract_slide_blocks(soup):
    for selector in ("section", "div.slide", "[data-slide]"):
        blocks = soup.select(selector)
        if blocks:
            return blocks
    return []


def _render_slide(slide, block, presentation_title, slide_width, slide_height,
                  slide_number, slide_count, footer_html, asset_base_path):
    container = block.find("body") or block
    background = _parse_color(container.get("style", ""), "background(?:-color)?")
    if background:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*background)
    title_el = container.find(["h1", "h2"])
    title_text = presentation_title or (title_el.get_text(" ", strip=True) if title_el else "")
    y = CONTENT_TOP
    if title_text:
        _add_title(slide, title_text, slide_width)
        y += TITLE_HEIGHT

    content_width = slide_width - 2 * MARGIN
    max_height = slide_height - y - MARGIN
    for element in container.children:
        if not getattr(element, "name", None):
            continue
        if element.name == "footer":
            continue
        if element.name in ("h1", "h2") and title_text:
            continue
        y = _add_element(slide, element, MARGIN, y, content_width, max_height, asset_base_path)
        if y > slide_height - MARGIN:
            raise PptxRenderError(f"conteúdo excede a altura após <{element.name}>.")
    footer = footer_html
    if footer is None:
        footer_el = container.find("footer")
        footer = footer_el.decode_contents() if footer_el else None
    if footer:
        footer = (footer.replace("{{slideNumber}}", str(slide_number))
                  .replace("{{slideCount}}", str(slide_count))
                  .replace("{slideNumber}", str(slide_number))
                  .replace("{slideCount}", str(slide_count)))
        _add_footer(slide, BeautifulSoup(footer, "html.parser").get_text(" ", strip=True), slide_width)


def _add_title(slide, text, slide_width):
    box = slide.shapes.add_textbox(
        Inches(MARGIN),
        Inches(0.35),
        Inches(slide_width - 2 * MARGIN),
        Inches(TITLE_HEIGHT),
    )
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)


def _add_element(slide, element, left, y, width, max_height, asset_base_path=None):
    name = element.name
    if name in ("h1", "h2", "h3", "h4", "h5", "h6", "p", "pre", "code", "span", "small", "label"):
        box = slide.shapes.add_textbox(Inches(left), Inches(y), Inches(width), Inches(0.6))
        paragraph = box.text_frame.paragraphs[0]
        _add_runs(paragraph, element, Pt(14 if name == "p" else max(18, 24 - (int(name[1]) - 1) * 2)))
        color = _parse_color(element.get("style", ""), "color")
        if color:
            paragraph.font.color.rgb = RGBColor(*color)
        return y + (0.4 if name == "p" else 0.55)

    if name in ("ul", "ol"):
        items = element.find_all("li", recursive=False)
        box = slide.shapes.add_textbox(Inches(left), Inches(y), Inches(width), Inches(0.4 * len(items) + 0.2))
        for index, item in enumerate(items):
            paragraph = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
            paragraph.text = ("• " if name == "ul" else f"{index + 1}. ") + item.get_text(" ", strip=True)
            paragraph.font.size = Pt(12)
        return y + 0.4 * len(items) + 0.3

    if name == "table":
        rows = element.find_all("tr")
        columns = max((len(row.find_all(["th", "td"])) for row in rows), default=0)
        if rows and columns:
            height = 0.4 * len(rows)
            if height > max_height:
                raise PptxRenderError(
                    f"tabela com {len(rows)} linhas não cabe na área restante."
                )
            table = slide.shapes.add_table(len(rows), columns, Inches(left), Inches(y), Inches(width), Inches(height)).table
            for row_index, row in enumerate(rows):
                for col_index, cell in enumerate(row.find_all(["th", "td"])):
                    if col_index >= columns:
                        continue
                    target = table.cell(row_index, col_index)
                    target.text = cell.get_text(" ", strip=True)
                    if cell.name == "th":
                        target.text_frame.paragraphs[0].font.bold = True
            return y + height + 0.2

    if name == "img":
        source = _image_source(element.get("src", ""), asset_base_path)
        if source:
            try:
                requested_width = float(element.get("width", 4))
            except (TypeError, ValueError):
                requested_width = 4
            image_width = min(width, max(0.5, requested_width))
            picture = slide.shapes.add_picture(source, Inches(left), Inches(y), width=Inches(image_width))
            return y + picture.height / 914400 + 0.15
        return y

    if name in ("div", "article", "main", "blockquote"):
        for child in element.children:
            if getattr(child, "name", None):
                y = _add_element(slide, child, left, y, width, max_height, asset_base_path)
                if y > CONTENT_TOP + max_height:
                    raise PptxRenderError(f"conteúdo dentro de <{name}> excede a altura disponível.")
        return y
    if element.get_text(" ", strip=True):
        raise PptxRenderError(f"elemento HTML <{name}> não é suportado pelo conversor PPTX.")
    return y


def _add_footer(slide, text, slide_width):
    box = slide.shapes.add_textbox(
        Inches(MARGIN), Inches(SLIDE_HEIGHT - 0.35),
        Inches(slide_width - 2 * MARGIN), Inches(0.2)
    )
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def _add_runs(paragraph, node, size):
    paragraph.font.size = size
    for child in node.children:
        if isinstance(child, str):
            if child.strip():
                run = paragraph.add_run()
                run.text = child
                run.font.size = size
            continue
        if not getattr(child, "name", None):
            continue
        run = paragraph.add_run()
        run.text = child.get_text(" ", strip=True)
        run.font.size = size
        run.font.bold = child.name in ("b", "strong")
        run.font.italic = child.name in ("i", "em")
        if child.name == "a":
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            if child.get("href"):
                run.hyperlink.address = child["href"]
        color = _parse_color(child.get("style", "")) if child.name == "span" else None
        if color:
            run.font.color.rgb = RGBColor(*color)


def _parse_color(style, property_name="color"):
    match = re.search(rf"{property_name}\s*:\s*#([0-9a-f]{{3}}|[0-9a-f]{{6}})\b", style, re.I)
    if not match:
        return None
    value = match.group(1)
    if len(value) == 3:
        value = "".join(char * 2 for char in value)
    return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))


def _apply_basic_css(soup, css):
    """Aplica regras CSS simples de tag, classe e id antes da conversão."""
    for selector_group, declarations in re.findall(r"([^{}]+)\{([^{}]*)\}", css or ""):
        style = "; ".join(part.strip() for part in declarations.split(";") if part.strip())
        for selector in selector_group.split(","):
            selector = selector.strip()
            if not selector or any(char in selector for char in (" ", ">", "+", "[", ":")):
                continue
            try:
                elements = soup.select(selector)
            except Exception:
                continue
            for element in elements:
                previous = element.get("style", "")
                element["style"] = f"{previous}; {style}" if previous else style


def _image_source(source, asset_base_path=None):
    source = source.strip()
    if source.startswith("data:image/"):
        match = re.match(r"data:image/[^;]+;base64,(.+)", source)
        return BytesIO(base64.b64decode(match.group(1))) if match else None
    if source.startswith(("http://", "https://")):
        raise ValueError("render_slide: imagens remotas não são suportadas; use arquivo local ou data URI.")
    try:
        candidate = source
        if not os.path.isabs(candidate):
            candidate = os.path.join(asset_base_path or os.getcwd(), candidate)
        return validate_path(candidate, True)
    except (OSError, ValueError) as exc:
        raise ValueError(f"render_slide: imagem inválida: {source}") from exc


def _validate_pptx(path):
    try:
        with zipfile.ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise ValueError("pacote ZIP interno corrompido")
        loaded = Presentation(path)
        if not loaded.slides:
            raise ValueError("apresentação sem slides")
    except Exception as exc:
        raise PptxRenderError(f"PPTX final inválido: {exc}") from exc
