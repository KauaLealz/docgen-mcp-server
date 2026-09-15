import os

import pytest
from pptx import Presentation

from services import read_doc_service, render_image_service, render_slide_service

pytestmark = pytest.mark.anyio


def _load_png_info(path):
    with open(path, "rb") as f:
        data = f.read()
    # PNG IHDR: width/height are bytes 16-24, big-endian; also check color
    # type byte (25) for alpha presence (6 = RGBA).
    assert data[:8] == b"\x89PNG\r\n\x1a\n"
    width = int.from_bytes(data[16:20], "big")
    height = int.from_bytes(data[20:24], "big")
    color_type = data[25]
    return width, height, color_type


async def test_render_image_full_page_png(tmp_path):
    out = str(tmp_path / "img.png")
    await render_image_service({
        "html": "<div style='width:100px;height:80px;background:#c00;'></div>",
        "css": "",
        "outputPath": out,
        "width": 100,
        "height": 80,
    })
    assert os.path.getsize(out) > 0
    w, h, _ = _load_png_info(out)
    assert (w, h) == (100, 80)


async def test_render_image_with_selector_crops_to_element(tmp_path):
    out = str(tmp_path / "badge.png")
    await render_image_service({
        "html": "<div class='badge' style='width:50px;height:50px;background:#0a0;'></div>",
        "css": "",
        "outputPath": out,
        "width": 300,
        "height": 300,
        "selector": ".badge",
    })
    w, h, _ = _load_png_info(out)
    assert (w, h) == (50, 50)


async def test_render_image_missing_selector_raises(tmp_path):
    out = str(tmp_path / "x.png")
    with pytest.raises(ValueError):
        await render_image_service({
            "html": "<div>no badge here</div>",
            "css": "",
            "outputPath": out,
            "width": 100,
            "height": 100,
            "selector": ".does-not-exist",
        })


async def test_render_image_transparent_background_has_alpha_channel(tmp_path):
    out = str(tmp_path / "transparent.png")
    await render_image_service({
        "html": "<div style='width:40px;height:40px;'></div>",
        "css": "",
        "outputPath": out,
        "width": 40,
        "height": 40,
        "transparentBackground": True,
    })
    _, _, color_type = _load_png_info(out)
    assert color_type == 6  # RGBA


async def test_render_image_scale_increases_pixel_dimensions(tmp_path):
    out = str(tmp_path / "scaled.png")
    await render_image_service({
        "html": "<div style='width:50px;height:50px;'></div>",
        "css": "",
        "outputPath": out,
        "width": 50,
        "height": 50,
        "scale": 2,
    })
    w, h, _ = _load_png_info(out)
    assert (w, h) == (100, 100)


async def test_render_image_invalid_format_raises(tmp_path):
    out = str(tmp_path / "x.gif")
    with pytest.raises(ValueError):
        await render_image_service({
            "html": "<div></div>", "css": "", "outputPath": out,
            "width": 10, "height": 10, "format": "gif",
        })


async def test_render_image_jpeg_format(tmp_path):
    out = str(tmp_path / "x.jpg")
    await render_image_service({
        "html": "<div style='width:20px;height:20px;background:#fff;'></div>",
        "css": "", "outputPath": out, "width": 20, "height": 20, "format": "jpeg",
    })
    with open(out, "rb") as f:
        header = f.read(3)
    assert header == b"\xff\xd8\xff"


async def test_render_slide_generates_pptx(tmp_path):
    out = str(tmp_path / "deck.pptx")
    result = await render_slide_service({
        "html": "<section><h1>Title</h1><p><strong>Hello</strong> world</p></section><section><h2>Second</h2><ul><li>Item</li></ul></section>",
        "css": "",
        "format": "pptx",
        "aspectRatio": "16:9",
        "outputPath": out,
        "title": "Deck",
        "author": "DocGen",
    })
    presentation = Presentation(out)
    assert len(presentation.slides) == 2
    assert presentation.core_properties.title == "Deck"
    assert presentation.core_properties.author == "DocGen"
    assert "PPTX" in result


async def test_render_slide_rejects_non_pptx_format(tmp_path):
    with pytest.raises(ValueError, match="format deve ser 'pptx'"):
        await render_slide_service({
            "html": "<section><h1>Title</h1></section>",
            "css": "",
            "format": "pdf",
            "aspectRatio": "16:9",
            "outputPath": str(tmp_path / "deck.pdf"),
        })


async def test_render_slide_rejects_pdf_only_options(tmp_path):
    with pytest.raises(ValueError, match="não suportados"):
        await render_slide_service({
            "html": "<section><p>Slide</p></section>",
            "css": "",
            "format": "pptx",
            "aspectRatio": "16:9",
            "outputPath": str(tmp_path / "deck.pptx"),
            "outline": True,
        })


async def test_render_slide_rejects_remote_images(tmp_path):
    with pytest.raises(ValueError, match="imagens remotas"):
        await render_slide_service({
            "html": "<section><img src='https://example.com/image.png'/></section>",
            "css": "",
            "format": "pptx",
            "aspectRatio": "16:9",
            "outputPath": str(tmp_path / "deck.pptx"),
        })


async def test_render_slide_supports_relative_images_and_correct_footer(tmp_path):
    from PIL import Image

    asset_dir = tmp_path / "assets"
    asset_dir.mkdir()
    Image.new("RGB", (20, 20), "red").save(asset_dir / "red.png")
    out = tmp_path / "relative.pptx"
    await render_slide_service({
        "html": "<section><p>One</p><img src='red.png'/></section><section><p>Two</p></section>",
        "css": "", "format": "pptx", "aspectRatio": "16:9", "outputPath": str(out),
        "assetBasePath": str(asset_dir), "footerHtml": "Página {slideNumber}/{slideCount}",
    })
    presentation = Presentation(out)
    assert len(presentation.slides) == 2
    for index, slide in enumerate(presentation.slides, start=1):
        texts = " ".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert f"Página {index}/2" in texts


async def test_read_doc_reads_pptx(tmp_path):
    out = tmp_path / "readable.pptx"
    await render_slide_service({
        "html": "<section><h1>Overview</h1><p>Content</p></section>",
        "css": "", "format": "pptx", "aspectRatio": "16:9", "outputPath": str(out),
    })
    result = read_doc_service(str(out), {})
    assert "Slide 1" in result["markdown"]
    assert "Overview" in result["markdown"]


async def test_render_slide_handles_dense_deck_with_more_than_seven_slides(tmp_path):
    slides = []
    for index in range(8):
        slides.append(
            f"<section><h2>Slide {index + 1}</h2><div class='compare'>"
            "<p>Left</p><p>Right</p><table><tr><th>A</th><th>B</th></tr>"
            "<tr><td>1</td><td>2</td></tr></table></div></section>"
        )
    out = tmp_path / "dense.pptx"
    await render_slide_service({
        "html": "".join(slides), "css": "", "format": "pptx",
        "aspectRatio": "16:9", "outputPath": str(out),
    })
    assert len(Presentation(out).slides) == 8


async def test_render_slide_reports_slide_number_for_unsupported_content(tmp_path):
    with pytest.raises(ValueError, match="slide 2/2.*<canvas>"):
        await render_slide_service({
            "html": "<section><p>OK</p></section><section><canvas>broken</canvas></section>",
            "css": "", "format": "pptx", "aspectRatio": "16:9",
            "outputPath": str(tmp_path / "invalid.pptx"),
        })


async def test_render_slide_applies_basic_css_and_validates_aspect_ratio(tmp_path):
    out = str(tmp_path / "styled.pptx")
    await render_slide_service({
        "html": "<section class='slide'><p>Styled</p></section>",
        "css": ".slide { background-color: #123456; } p { color: #ff0000; }",
        "format": "pptx",
        "aspectRatio": "4:3",
        "outputPath": out,
    })
    presentation = Presentation(out)
    assert presentation.slide_width > presentation.slide_height
    with pytest.raises(ValueError, match="aspectRatio"):
        await render_slide_service({
            "html": "<section><p>Invalid</p></section>",
            "css": "",
            "format": "pptx",
            "aspectRatio": "1:1",
            "outputPath": str(tmp_path / "invalid.pptx"),
        })
