import json
import os

import docx

from services import write_docx_from_content


def _cell_texts(table):
    return [[cell.text for cell in row.cells] for row in table.rows]


def test_ragged_table_rows_do_not_lose_or_shift_cells(tmp_path):
    content = json.dumps(
        [
            {
                "type": "table",
                "headers": ["A", "B"],
                "rows": [
                    ["1", "2", "3"],  # more cells than the header
                    ["only-one"],  # fewer cells than the header
                ],
            }
        ]
    )
    out = str(tmp_path / "ragged.docx")
    write_docx_from_content(content, "json", out)

    doc = docx.Document(out)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    texts = _cell_texts(table)
    assert texts[0] == ["A", "B", ""]
    assert texts[1] == ["1", "2", "3"]
    assert texts[2] == ["only-one", "", ""]


def test_unknown_html_tag_is_not_silently_dropped(tmp_path):
    from services import html_to_docx_paragraphs

    doc = docx.Document()
    html_to_docx_paragraphs(doc, "<section><p>inside section</p></section><foo>bare text</foo>")
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "inside section" in all_text
    assert "bare text" in all_text


def test_table_cell_with_multiple_paragraphs_renders_as_multiple_paragraphs(tmp_path):
    content = json.dumps(
        [
            {
                "type": "table",
                "rows": [
                    [
                        [
                            {"type": "paragraph", "text": "first"},
                            {"type": "paragraph", "text": "second"},
                        ],
                        "plain",
                    ]
                ],
            }
        ]
    )
    out = str(tmp_path / "multiparacell.docx")
    write_docx_from_content(content, "json", out)

    doc = docx.Document(out)
    table = doc.tables[0]
    cell = table.rows[0].cells[0]
    paragraph_texts = [p.text for p in cell.paragraphs if p.text]
    assert paragraph_texts == ["first", "second"]
    assert table.rows[0].cells[1].text == "plain"


def test_markdown_table_still_works_end_to_end(tmp_path):
    md = "# Report\n\n| Name | Score |\n| --- | --- |\n| Alice | 10 |\n| Bob | 20 |\n"
    out = str(tmp_path / "md.docx")
    write_docx_from_content(md, "markdown", out)

    doc = docx.Document(out)
    assert doc.paragraphs[0].text == "Report"
    assert len(doc.tables) == 1
    assert _cell_texts(doc.tables[0]) == [["Name", "Score"], ["Alice", "10"], ["Bob", "20"]]


def test_invalid_content_format_raises():
    import pytest

    with pytest.raises(ValueError):
        write_docx_from_content("hello", "plain", "/tmp/should-not-be-created.docx")
